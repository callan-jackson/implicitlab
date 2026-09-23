"""Client deliverables: the workbook and the deck.

The property that matters most is that a deliverable cannot disagree with the
dashboard. Both exports are built from ``analyse_cohort``, and these tests
check that the numbers that land in the files are those numbers — plus the
structural guarantees a research manager relies on: every tab present, every
recruited participant accounted for, every raw trial in the audit trail.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.cohort.analysis import analyse_cohort
from app.cohort.export_pptx import build_deck
from app.cohort.export_xlsx import RAW_SHEET, SHEETS, build_workbook
from app.cohort.registry import Batch, _compact
from app.cohort.scoring import CohortParams, prepare
from app.cohort.simulate import simulate_batch


@pytest.fixture(scope="module")
def batch() -> Batch:
    sim = simulate_batch(80, seed=3)
    trials = _compact(sim.trials)
    prep = prepare(trials, list(sim.meta["segments"]))
    return Batch("test", {**sim.meta, "id": "test"}, {"rows_read": len(trials)}, trials, prep,
                 datetime.now(timezone.utc).isoformat())


@pytest.fixture(scope="module")
def workbook(batch):
    raw = build_workbook(batch, CohortParams(), segment_by="exposure", method="permutation")
    return load_workbook(io.BytesIO(raw), read_only=True)


@pytest.fixture(scope="module")
def deck(batch):
    return Presentation(io.BytesIO(build_deck(batch, CohortParams(), segment_by="exposure")))


# ---------------------------------------------------------------- workbook


def test_workbook_has_the_five_tabs_in_order(workbook):
    assert workbook.sheetnames == SHEETS


def test_summary_d_matches_the_analysis(batch, workbook):
    report = analyse_cohort(batch.prep, batch.meta, CohortParams(), segment_by="exposure")
    ws = workbook["Executive Summary"]
    header_row = next(r for r in ws.iter_rows() if r[0].value == "Attribute")
    cols = [c.value for c in header_row]
    d_col = cols.index("Mean D")
    sig_col = next(i for i, c in enumerate(cols) if c and c.startswith("Significant"))
    rows = {}
    for r in ws.iter_rows(min_row=header_row[0].row + 1, max_row=header_row[0].row + 3):
        rows[r[0].value] = r
    for attr in report["attributes"]:
        row = rows[attr["pole_a"]]
        assert row[d_col].value == pytest.approx(attr["overall"]["mean_d"], abs=1e-12)
        assert row[sig_col].value == ("Yes" if attr["overall"]["permutation"]["sig"] else "No")


def test_every_recruited_participant_is_in_the_participant_tab(batch, workbook):
    ws = workbook["Participant D-scores"]
    assert ws.max_row - 1 == batch.prep.n_participants
    header = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert header[:2] == ["participant_id", "status"]
    statuses = [r[1].value for r in ws.iter_rows(min_row=2)]
    # Included first, then exclusions, each labelled with its reason.
    first_excluded = next(i for i, s in enumerate(statuses) if s != "Included")
    assert all(s == "Included" for s in statuses[:first_excluded])
    assert all(s.startswith("Excluded — ") for s in statuses[first_excluded:])


def test_raw_trial_log_is_complete_and_exact(batch, workbook):
    ws = workbook[RAW_SHEET]
    rows = list(ws.iter_rows(values_only=True))
    assert len(rows) - 1 == len(batch.trials)
    header = rows[0]
    assert header[0] == "participant_id"
    lat = header.index("latency_ms")
    pid = header.index("participant_id")
    ok = header.index("correct")
    for i in (0, len(batch.trials) // 2, len(batch.trials) - 1):
        src = batch.trials.iloc[i]
        assert rows[i + 1][pid] == src["participant_id"]
        assert rows[i + 1][lat] == pytest.approx(float(src["latency_ms"]))
        assert rows[i + 1][ok] is bool(src["correct"])


def test_parameters_are_recorded(batch):
    params = CohortParams(lower_ms=400, min_accuracy=0.8)
    wb = load_workbook(io.BytesIO(build_workbook(batch, params, method="parametric")),
                       read_only=True)
    ws = wb["Method & Parameters"]
    settings = {r[0].value: r[1].value for r in ws.iter_rows(min_row=5) if r[0].value}
    assert settings["lower_ms"] == 400
    assert settings["min_accuracy"] == 0.8
    assert settings["Parameters are defaults"] == "No"
    assert settings["Significance method"].startswith("t-tests")


def test_segment_tab_covers_every_variable(batch, workbook):
    ws = workbook["Segment Breakdown"]
    titles = {c.value for (c,) in ws.iter_rows(max_col=1) if c.value}
    for spec in batch.meta["segments"].values():
        assert any(str(t).startswith(spec["label"]) for t in titles)


# ---------------------------------------------------------------- deck


def test_deck_has_nine_slides_with_charts(deck):
    assert len(deck.slides) == 9
    pictures = sum(1 for s in deck.slides for sh in s.shapes
                   if sh.shape_type == MSO_SHAPE_TYPE.PICTURE)
    assert pictures >= 5


def test_no_empty_text_boxes(deck):
    for i, slide in enumerate(deck.slides, 1):
        for sh in slide.shapes:
            if sh.shape_type == MSO_SHAPE_TYPE.TEXT_BOX:
                assert sh.text_frame.text.strip(), f"empty text box on slide {i}"


def test_deck_carries_the_takeaways_and_the_synthetic_flag(batch, deck):
    report = analyse_cohort(batch.prep, batch.meta, CohortParams(), segment_by="exposure")
    text = "\n".join(sh.text_frame.text for sh in deck.slides[1].shapes if sh.has_text_frame)
    for t in report["takeaways"]["permutation"]:
        assert t["headline"] in text
    for slide in deck.slides:
        assert any("Synthetic data" in sh.text_frame.text
                   for sh in slide.shapes if sh.has_text_frame)


# ---------------------------------------------------------------- HTTP


def test_export_endpoints():
    from fastapi.testclient import TestClient

    from app.main import app

    with TestClient(app) as c:
        r = c.get("/api/cohort/batches/demo/export.pptx?segment_by=loyalty&method=parametric")
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats-officedocument.presentationml")
        assert "attachment" in r.headers["content-disposition"]
        assert len(Presentation(io.BytesIO(r.content)).slides) == 9
        assert c.get("/api/cohort/batches/demo/export.xlsx?method=bogus").status_code == 400
        assert c.get("/api/cohort/batches/nope/export.xlsx").status_code == 404
