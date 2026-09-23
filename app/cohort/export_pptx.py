"""The client deck: nine slides, generated, never hand-edited.

A research manager's deliverable is a slide deck, and the slow, error-prone
part of producing one is copying numbers from an analysis output into slides.
This builds the deck from the same :func:`analyse_cohort` report the Studio and
the workbook use, with the same parameters, so a figure on a slide is by
construction the figure in the workbook.

The slide titles are *findings*, not topics ("Aurelia owns Premium", not
"Heatmap") — the convention in agency decks, where a client skimming titles
alone should get the story. Every title is generated from the same rules that
write the executive takeaways, so a title can never claim more than the
statistics under it support.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from ..config import get_settings
from . import charts
from .analysis import analyse_cohort
from .scoring import CohortParams

W, H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.6)
FONT = "Arial"

INK = RGBColor(0x0B, 0x0B, 0x0B)
INK_2 = RGBColor(0x52, 0x51, 0x4E)
MUTED = RGBColor(0x89, 0x87, 0x81)
ACCENT = RGBColor(0x4A, 0x3A, 0xA7)
POS = RGBColor(0x2A, 0x78, 0xD6)
NEG = RGBColor(0xEB, 0x68, 0x34)
GOOD = RGBColor(0x00, 0x63, 0x00)
WARN = RGBColor(0xB4, 0x53, 0x09)
RULE = RGBColor(0xE1, 0xE0, 0xD9)
PANEL = RGBColor(0xF6, 0xF5, 0xF1)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

METHOD_LABEL = {
    "permutation": "permutation tests on D (distribution-free)",
    "parametric": "t-tests on log-transformed reaction times",
}


# --------------------------------------------------------------------------
# primitives
# --------------------------------------------------------------------------


def _brand(s: str | None) -> str:
    return s.title() if s and s.isupper() else (s or "")


def _fd(v: float | None, dp: int = 2) -> str:
    return "—" if v is None else f"{v:+.{dp}f}"


def _fq(q: float | None) -> str:
    if q is None:
        return "—"
    return "<.001" if q < 0.001 else f"{q:.3f}".lstrip("0")


def _q(q: float | None) -> str:
    """'q < .001' or 'q = .012' — never 'q = <.001'."""
    v = _fq(q)
    return f"q {v[0]} {v[1:]}" if v.startswith("<") else f"q = {v}"


def _text(slide, x, y, w, h, text: str = "", *, size: float = 14, bold: bool = False,
          color: RGBColor = INK, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
          italic: bool = False):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    _style(r, size=size, bold=bold, color=color, italic=italic)
    return box


def _style(run, *, size: float, bold: bool = False, color: RGBColor = INK, italic: bool = False):
    f = run.font
    f.name = FONT
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.color.rgb = color


def _para(tf, text: str, *, size: float, bold: bool = False, color: RGBColor = INK,
          space_before: float = 0, first: bool = False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.space_before = Pt(space_before)
    r = p.add_run()
    r.text = text
    _style(r, size=size, bold=bold, color=color)
    return p


def _rect(slide, x, y, w, h, fill: RGBColor):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    s.line.fill.background()
    s.shadow.inherit = False
    return s


def _picture(slide, png: bytes, x, y, *, max_w, max_h):
    """Place an image inside a box, preserving aspect ratio, centred."""
    from PIL import Image

    im = Image.open(io.BytesIO(png))
    ratio = im.width / im.height
    w, h = max_w, int(max_w / ratio)
    if h > max_h:
        h, w = max_h, int(max_h * ratio)
    left = x + (max_w - w) // 2
    return slide.shapes.add_picture(io.BytesIO(png), left, y, Emu(w), Emu(h))


class Deck:
    def __init__(self, study: str, simulated: bool):
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = W, H
        self.layout = self.prs.slide_layouts[6]  # blank
        self.study = study
        self.simulated = simulated
        self.n = 0

    def slide(self, title: str | None, kicker: str | None = None):
        s = self.prs.slides.add_slide(self.layout)
        self.n += 1
        if title is not None:
            _rect(s, MARGIN, Inches(0.55), Inches(0.08), Inches(0.62), ACCENT)
            if kicker:
                _text(s, MARGIN + Inches(0.25), Inches(0.42), Inches(11.5), Inches(0.3),
                      kicker.upper(), size=10.5, bold=True, color=ACCENT)
            # Titles are generated, so their length is not known in advance;
            # long ones step down a size rather than wrap into the body.
            size = 26 if len(title) <= 62 else 22 if len(title) <= 76 else 19
            _text(s, MARGIN + Inches(0.25), Inches(0.7), Inches(11.9), Inches(0.6),
                  title, size=size, bold=True, color=INK)
        # Footer.
        y = H - Inches(0.45)
        _rect(s, MARGIN, y - Inches(0.08), W - 2 * MARGIN, Emu(9525), RULE)
        _text(s, MARGIN, y, Inches(5.5), Inches(0.3), f"ImplicitLab · {self.study}",
              size=9, color=MUTED)
        if self.simulated:
            _text(s, Inches(4.9), y, Inches(3.6), Inches(0.3),
                  "Synthetic data — demonstration only", size=9, bold=True, color=WARN,
                  align=PP_ALIGN.CENTER)
        _text(s, W - MARGIN - Inches(1.0), y, Inches(1.0), Inches(0.3), str(self.n),
              size=9, color=MUTED, align=PP_ALIGN.RIGHT)
        return s

    def bytes(self) -> bytes:
        buf = io.BytesIO()
        self.prs.save(buf)
        return buf.getvalue()


BODY_TOP = Inches(1.55)
BODY_H = H - BODY_TOP - Inches(0.75)
BODY_W = W - 2 * MARGIN


# --------------------------------------------------------------------------
# slides
# --------------------------------------------------------------------------


def _title_slide(deck: Deck, batch, report: dict, method: str, generated: str) -> None:
    s = deck.slide(None)
    _rect(s, 0, 0, Inches(0.35), H, ACCENT)
    sample = report["sample"]
    _text(s, Inches(1.1), Inches(1.6), Inches(11), Inches(0.4),
          "IMPLICIT ASSOCIATION STUDY — TOPLINE", size=13, bold=True, color=ACCENT)
    _text(s, Inches(1.1), Inches(2.1), Inches(11.2), Inches(1.6),
          batch.meta.get("name", "Cohort study"), size=40, bold=True, color=INK)
    attrs = ", ".join(a.get("pole_a") or a["key"] for a in report["attributes"])
    _text(s, Inches(1.1), Inches(3.75), Inches(11), Inches(0.9),
          f"{sample['included']:,} respondents analysed of {sample['recruited']:,} recruited · "
          f"attributes tested: {attrs}", size=18, color=INK_2)
    _text(s, Inches(1.1), Inches(4.5), Inches(11), Inches(0.5),
          f"Significance: {METHOD_LABEL[method]}, Benjamini–Hochberg corrected · {generated}",
          size=13, color=MUTED)


def _takeaways_slide(deck: Deck, report: dict, method: str) -> None:
    s = deck.slide("Executive takeaways", "Summary")
    items = report["takeaways"][method]
    gap = Inches(0.08)
    row_h = min(Inches(0.92), int((BODY_H - gap * len(items)) / max(1, len(items))))
    y = BODY_TOP
    for t in items:
        # Accent, not the blue/orange direction colours: a takeaway can favour
        # either brand, and blue means "leans target A" on every chart.
        colour = {"positive": ACCENT, "neutral": MUTED}.get(t.get("tone"), MUTED)
        if t["kind"] == "sample":
            colour = INK_2
        _rect(s, MARGIN, y + Inches(0.06), Inches(0.1), row_h - Inches(0.14), colour)
        box = s.shapes.add_textbox(MARGIN + Inches(0.3), y, BODY_W - Inches(0.3), row_h)
        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        _para(tf, t["headline"], size=15, bold=True, first=True)
        _para(tf, t["detail"], size=11.5, color=INK_2, space_before=2)
        y += row_h + gap


def _sample_slide(deck: Deck, report: dict) -> None:
    sample = report["sample"]
    s = deck.slide(f"{sample['included']:,} of {sample['recruited']:,} respondents met every "
                   f"quality criterion", "Sample & data quality")
    _picture(s, charts.funnel(report, width=7.4, height=4.3), MARGIN, BODY_TOP,
             max_w=Inches(7.6), max_h=BODY_H)
    x = MARGIN + Inches(7.9)
    w = BODY_W - Inches(7.9)
    _rect(s, x, BODY_TOP, w, BODY_H, PANEL)
    box = s.shapes.add_textbox(x + Inches(0.25), BODY_TOP + Inches(0.2), w - Inches(0.5),
                               BODY_H - Inches(0.4))
    tf = box.text_frame
    tf.word_wrap = True
    _para(tf, "Exclusion rules, applied in order", size=13, bold=True, first=True)
    for st in report["funnel"]:
        if st["key"] in ("recruited", "final"):
            continue
        _para(tf, f"{st['label']} — {st['removed']:,} removed", size=11.5, bold=True,
              color=INK, space_before=7)
        _para(tf, st["reason"], size=10, color=INK_2)
    tr = report["trimming"]
    p = report["params"]
    lower = "no lower trim" if not p["lower_ms"] else f"trials under {p['lower_ms']:.0f} ms"
    _para(tf, "Trial-level trimming", size=11.5, bold=True, space_before=10)
    _para(tf, f"{tr['removed_upper'] + tr['removed_lower']:,} of {tr['combined_trials']:,} "
              f"scored trials removed ({tr['removed_pct']:.1%}): above {p['upper_ms']:,.0f} ms, "
              f"{lower}. Participants are screened before trimming, so a trim cannot hide a "
              f"fast responder.", size=10, color=INK_2)


def _heatmap_slide(deck: Deck, report: dict, method: str) -> None:
    heads = [t["headline"] for t in report["takeaways"][method] if t["kind"] == "association"]
    title = "; ".join(heads[:2]) if heads else "Association strength by attribute"
    s = deck.slide(title, "Association map")
    _picture(s, charts.heatmap(report, method, width=10.5), MARGIN, BODY_TOP,
             max_w=BODY_W, max_h=BODY_H - Inches(0.6))
    a0 = report["attributes"][0] if report["attributes"] else {}
    _text(s, MARGIN, H - Inches(1.2), BODY_W, Inches(0.45),
          f"Mean D per cell. D is relative: positive means respondents paired "
          f"{_brand(a0.get('target_a'))} with the attribute faster than "
          f"{_brand(a0.get('target_b'))}. Filled dot = significant after FDR correction.",
          size=11, color=INK_2)


def _forest_slide(deck: Deck, report: dict, method: str, seg_label: str) -> None:
    s = deck.slide(f"Where each segment sits — by {seg_label.lower()}", "Effect sizes")
    _picture(s, charts.forest(report, method, width=8.2), MARGIN, BODY_TOP,
             max_w=Inches(8.6), max_h=BODY_H)
    x = MARGIN + Inches(8.9)
    w = BODY_W - Inches(8.9)
    box = s.shapes.add_textbox(x, BODY_TOP + Inches(0.2), w, BODY_H - Inches(0.4))
    tf = box.text_frame
    tf.word_wrap = True
    _para(tf, "How to read it", size=13, bold=True, first=True)
    for line in (
        "Each dot is a mean D; the line is its 95% bootstrap confidence interval, resampled "
        "over participants.",
        "A line that crosses zero is not distinguishable from no association.",
        "Filled = significant after Benjamini–Hochberg correction; hollow = not.",
        "Blue leans towards the first brand, orange towards the second.",
        "Cells under 30 respondents are marked low base — indicative only.",
    ):
        _para(tf, line, size=11, color=INK_2, space_before=8)


def _story_slide(deck: Deck, report: dict, method: str, seg_label: str) -> None:
    comps = report["comparisons"]
    sig = sorted((c for c in comps if c[method].get("sig")),
                 key=lambda c: -abs(c.get("hedges_g") or 0))
    attrs = {a["key"]: a for a in report["attributes"]}
    if not sig:
        s = deck.slide(f"No reliable differences by {seg_label.lower()}", "Segment story")
        biggest = max(comps, key=lambda c: abs(c.get("diff") or 0)) if comps else None
        msg = (f"None of the {len(comps)} {seg_label.lower()} comparisons survives correction "
               f"for multiple comparisons.")
        if biggest:
            pole = attrs[biggest["attribute"]].get("pole_a")
            msg += (f" The largest gap — {pole}, {biggest['level_a']} vs {biggest['level_b']} "
                    f"({_fd(biggest['diff'])}, {_q(biggest[method].get('q'))}) — is within "
                    f"what chance produces across this many tests.")
        _text(s, MARGIN, BODY_TOP + Inches(0.1), Inches(11.5), Inches(1.2), msg, size=16, color=INK_2)
        # With no segment story to tell, show how much respondents vary
        # instead: the spread is why a segment gap has to be large to register.
        _text(s, MARGIN, BODY_TOP + Inches(1.35), BODY_W, Inches(0.35),
              "Respondent-level D distribution per attribute", size=12, bold=True)
        _picture(s, charts.d_histograms(report, width=11.5, height=2.9), MARGIN,
                 BODY_TOP + Inches(1.75), max_w=BODY_W, max_h=BODY_H - Inches(1.8))
        return
    c = sig[0]
    attr = attrs[c["attribute"]]
    pole = attr.get("pole_a") or c["attribute"]
    a_brand = _brand(attr.get("target_a"))
    hi, lo = (c["level_a"], c["level_b"]) if (c["diff"] or 0) > 0 else (c["level_b"], c["level_a"])
    s = deck.slide(f"{seg_label} moves {pole}", f"Segment story · stronger {a_brand}–{pole} "
                   f"link among {hi} than {lo}")
    # Two big-number callouts.
    col_w = Inches(3.7)
    for i, (lvl, mean, n) in enumerate(((c["level_a"], c["mean_a"], c["n_a"]),
                                        (c["level_b"], c["mean_b"], c["n_b"]))):
        x = MARGIN + i * (col_w + Inches(0.3))
        _rect(s, x, BODY_TOP, col_w, Inches(2.3), PANEL)
        _text(s, x + Inches(0.3), BODY_TOP + Inches(0.2), col_w - Inches(0.6), Inches(0.4),
              f"{lvl}  ·  n = {n}", size=13, bold=True, color=INK_2)
        _text(s, x + Inches(0.3), BODY_TOP + Inches(0.6), col_w - Inches(0.6), Inches(1.1),
              _fd(mean), size=54, bold=True, color=POS if (mean or 0) >= 0 else NEG)
        _text(s, x + Inches(0.3), BODY_TOP + Inches(1.75), col_w - Inches(0.6), Inches(0.4),
              f"mean D, {pole}", size=11, color=MUTED)
    x = MARGIN
    y = BODY_TOP + Inches(2.6)
    box = s.shapes.add_textbox(x, y, Inches(7.7), BODY_H - Inches(2.7))
    tf = box.text_frame
    tf.word_wrap = True
    _para(tf, f"Difference {_fd(c['diff'])}  (95% CI {_fd(c['ci_low'])} to {_fd(c['ci_high'])})",
          size=16, bold=True, first=True)
    _para(tf, f"Hedges' g = {c['hedges_g']:.2f} · permutation {_q(c['permutation'].get('q'))} · "
              f"Welch t on log RT {_q(c['parametric'].get('q'))}"
              + (" · both tests agree" if c["agree"] else " · the two tests disagree — borderline"),
          size=12, color=INK_2, space_before=6)
    if c["low_base"]:
        _para(tf, "Low base in at least one group — indicative only.", size=12, bold=True,
              color=WARN, space_before=6)
    g = c.get("hedges_g")
    if g is not None:
        _para(tf, f"Reading it: {hi} respondents pair {a_brand} with {pole} measurably faster "
                  f"than {lo} respondents — a gap of {abs(c['diff']):.2f} D units, or "
                  f"{abs(g):.2f} standard deviations of respondent-level D.",
              size=12, color=INK_2, space_before=12)
    seg_note = (report["batch"].get("segments", {}).get(report.get("segment_by")) or {}).get("note")
    if seg_note:
        _para(tf, f"About this variable: {seg_note}", size=11, color=MUTED, space_before=8)
    others = [o for o in sig[1:3]]
    if others:
        _para(tf, "Also significant:", size=12, bold=True, space_before=12)
        for o in others:
            _para(tf, f"{attrs[o['attribute']].get('pole_a')}: {o['level_a']} {_fd(o['mean_a'])} vs "
                      f"{o['level_b']} {_fd(o['mean_b'])} ({_q(o[method].get('q'))})",
                  size=11.5, color=INK_2, space_before=3)
    # Right: the forest for just that attribute.
    sub = {**report, "attributes": [attr]}
    _picture(s, charts.forest(sub, method, width=4.4, height=4.2), MARGIN + Inches(7.9),
             BODY_TOP, max_w=BODY_W - Inches(7.9), max_h=BODY_H)


def _distribution_slide(deck: Deck, report: dict) -> None:
    attr = report["attributes"][0]
    shape = report["distributions"][attr["key"]]["shape"]
    con, inc = shape.get("congruent", {}), shape.get("incongruent", {})
    s = deck.slide("Reaction times are skewed — so the tests don't assume normality",
                   "Why these statistics")
    _picture(s, charts.rt_distribution(report, attr["key"], width=9.6, height=3.6),
             MARGIN, BODY_TOP, max_w=BODY_W, max_h=Inches(3.5))
    y = BODY_TOP + Inches(3.7)
    rows = [
        ("", "Raw RT", "Log RT", "A normal curve"),
        ("Skewness", f"{con.get('skew', 0):.2f} / {inc.get('skew', 0):.2f}",
         f"{con.get('log_skew', 0):.2f} / {inc.get('log_skew', 0):.2f}", "0"),
        ("Excess kurtosis", f"{con.get('excess_kurtosis', 0):.2f} / {inc.get('excess_kurtosis', 0):.2f}",
         f"{con.get('log_excess_kurtosis', 0):.2f} / {inc.get('log_excess_kurtosis', 0):.2f}", "0"),
        ("Trials beyond mean + 2 SD", f"{con.get('beyond_2sd', 0):.1%} / {inc.get('beyond_2sd', 0):.1%}",
         "—", "2.3%"),
    ]
    tbl = s.shapes.add_table(len(rows), 4, MARGIN, y, Inches(7.4), Inches(1.5)).table
    widths = [Inches(2.6), Inches(1.7), Inches(1.7), Inches(1.4)]
    for j, w in enumerate(widths):
        tbl.columns[j].width = w
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = ACCENT if i == 0 else (PANEL if i % 2 else WHITE)
            tf = cell.text_frame
            tf.paragraphs[0].text = ""
            run = tf.paragraphs[0].add_run()
            run.text = v
            _style(run, size=11, bold=i == 0 or j == 0, color=WHITE if i == 0 else INK)
            cell.margin_top = cell.margin_bottom = Inches(0.03)
    _text(s, MARGIN + Inches(7.8), y, BODY_W - Inches(7.8), Inches(1.7),
          f"{attr.get('pole_a')} task, correct trials, congruent / incongruent. Raw latencies "
          f"have a long right tail a t-test would mis-model; logs are close to symmetric. So the "
          f"parametric tests run on log RTs, and the permutation tests on D assume no shape "
          f"at all. Both are reported; slide {deck.n + 1} shows they agree.",
          size=11.5, color=INK_2)


def _robustness_slide(deck: Deck, report: dict, method: str, seg_label: str) -> None:
    ag = report["agreement"]
    title = ("Every conclusion survives a change of statistical test" if ag["n_agree"] == ag["n_tests"]
             else f"{ag['n_agree']} of {ag['n_tests']} significance calls hold under both tests")
    s = deck.slide(title, "Robustness")
    rows = []
    for a in report["attributes"]:
        for c in [a["overall"]] + a["levels"]:
            rows.append((f"{a.get('pole_a')} — {c['level'] or 'all respondents'}", c))
    for c in report["comparisons"]:
        pole = next(a.get("pole_a") for a in report["attributes"] if a["key"] == c["attribute"])
        rows.append((f"{pole} — {c['level_a']} vs {c['level_b']}", c))
    rows.sort(key=lambda r: r[1]["agree"])  # disagreements first
    max_rows = 14
    shown = rows[:max_rows]
    header = ["Test", "t-test (log RT) q", "Sig.", "Permutation q", "Sig.", "Agree"]
    tbl = s.shapes.add_table(len(shown) + 1, len(header), MARGIN, BODY_TOP,
                             Inches(8.6), Inches(0.3) * (len(shown) + 1)).table
    for j, w in enumerate([Inches(3.8), Inches(1.3), Inches(0.7), Inches(1.3), Inches(0.7), Inches(0.8)]):
        tbl.columns[j].width = w
    for j, h in enumerate(header):
        cell = tbl.cell(0, j)
        cell.fill.solid()
        cell.fill.fore_color.rgb = ACCENT
        run = cell.text_frame.paragraphs[0].add_run()
        run.text = h
        _style(run, size=10, bold=True, color=WHITE)
    for i, (label, c) in enumerate(shown, 1):
        vals = [label, _fq(c["parametric"].get("q")), "●" if c["parametric"].get("sig") else "○",
                _fq(c["permutation"].get("q")), "●" if c["permutation"].get("sig") else "○",
                "Yes" if c["agree"] else "No"]
        for j, v in enumerate(vals):
            cell = tbl.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PANEL if i % 2 else WHITE
            cell.margin_top = cell.margin_bottom = Inches(0.02)
            run = cell.text_frame.paragraphs[0].add_run()
            run.text = v
            _style(run, size=10, color=(WARN if (j == 5 and not c["agree"]) else INK),
                   bold=(j == 5 and not c["agree"]))
            if j:
                cell.text_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        tbl.rows[i].height = Inches(0.3)
    x = MARGIN + Inches(8.9)
    box = s.shapes.add_textbox(x, BODY_TOP, BODY_W - Inches(8.9), BODY_H)
    tf = box.text_frame
    tf.word_wrap = True
    _para(tf, f"{ag['n_agree']} / {ag['n_tests']}", size=40, bold=True,
          color=GOOD if ag["n_agree"] == ag["n_tests"] else WARN, first=True)
    _para(tf, "significance calls identical under both tests", size=12, color=INK_2)
    _para(tf, "Parametric: t-tests on each respondent's mean log-RT difference; Welch's test "
              "between segments. Non-parametric: sign-flip and label-shuffle permutation tests "
              "on D. q-values are Benjamini–Hochberg corrected within each family.",
          size=10.5, color=INK_2, space_before=14)
    if len(rows) > max_rows:
        _para(tf, f"Showing {max_rows} of {len(rows)} tests, disagreements first. The full "
                  f"table is in the workbook's Segment Breakdown tab.", size=10, color=MUTED,
              space_before=10)


def _method_slide(deck: Deck, batch, report: dict, params: CohortParams, method: str) -> None:
    s = deck.slide("Method & caveats", "Appendix")
    p = params.to_dict()
    left = [
        ("Measure", "Greenwald, Nosek & Banaji (2003) improved D per respondent per attribute; "
                    "built-in error penalty (forced correction), inclusive SD, practice and test "
                    "quotients averaged."),
        ("Exclusions", f"Participant-level, in order: incomplete data; more than "
                       f"{p['fast_limit']:.0%} of trials under {p['fast_ms']:.0f} ms; accuracy "
                       f"under {p['min_accuracy']:.0%}; fewer than {p['min_trials']} trials per "
                       f"cell after trimming. Complete cases only."),
        ("Trimming", f"Trials above {p['upper_ms']:,.0f} ms removed; "
                     + ("no lower trim." if not p["lower_ms"] else f"trials under {p['lower_ms']:.0f} ms removed.")
                     + ("" if params.is_default() else " Non-default settings chosen by the analyst.")),
        ("Inference", "95% percentile bootstrap CIs over participants (2,000 resamples). "
                      "Flags use " + METHOD_LABEL[method] + "; both methods are in the workbook."),
        ("Multiplicity", "Benjamini–Hochberg false-discovery-rate correction within each family "
                         "(associations; segment differences). Significant = q < .05."),
    ]
    right = [
        ("D is relative", "A two-target IAT compares brands. It supports 'more associated than'; "
                          "it cannot say a brand is premium in absolute terms."),
        ("Group-level measure", "Individual D-scores are noisy (test–retest around r = .5). "
                                "Read cells, not respondents."),
        ("Size bands", "0.15 / 0.35 / 0.65 are the Project Implicit reporting convention, not "
                       "Cohen's d."),
        ("Low base", "Cells under 30 respondents are marked low base — indicative only."),
        ("Reproducibility", f"Batch {batch.id}; every figure is a pure function of the batch and "
                            f"these parameters. ImplicitLab v{get_settings().version}."),
    ]
    if batch.meta.get("source") == "simulated":
        right.append(("Synthetic data", "Generated with planted effects for demonstration "
                                        f"(seed {batch.meta.get('seed')}). Nobody took these tests."))
    col_w = (BODY_W - Inches(0.5)) / 2
    for k, items in enumerate((left, right)):
        box = s.shapes.add_textbox(MARGIN + k * (col_w + Inches(0.5)), BODY_TOP, int(col_w), BODY_H)
        tf = box.text_frame
        tf.word_wrap = True
        for i, (h, body) in enumerate(items):
            _para(tf, h, size=12, bold=True, color=ACCENT, first=i == 0, space_before=0 if i == 0 else 9)
            _para(tf, body, size=10.5, color=INK_2, space_before=1)


# --------------------------------------------------------------------------
# entry point
# --------------------------------------------------------------------------


def build_deck(batch, params: CohortParams, *, segment_by: str | None = None,
               method: str = "permutation") -> bytes:
    report = analyse_cohort(batch.prep, batch.meta, params, segment_by=segment_by)
    seg = report["segment_by"]
    seg_label = (batch.meta.get("segments", {}).get(seg) or {}).get("label", seg or "segment")
    generated = datetime.now(timezone.utc).strftime("%d %B %Y")
    deck = Deck(batch.meta.get("name", "Cohort study"), batch.meta.get("source") == "simulated")

    _title_slide(deck, batch, report, method, generated)
    _takeaways_slide(deck, report, method)
    _sample_slide(deck, report)
    _heatmap_slide(deck, report, method)
    _forest_slide(deck, report, method, seg_label)
    _story_slide(deck, report, method, seg_label)
    _distribution_slide(deck, report)
    _robustness_slide(deck, report, method, seg_label)
    _method_slide(deck, batch, report, params, method)
    return deck.bytes()
