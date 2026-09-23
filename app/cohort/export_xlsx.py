"""The analyst's workbook: five tabs, one source of numbers.

Built with pandas + openpyxl from the same :func:`analyse_cohort` report the
Studio renders, with the same parameters, so the workbook a research manager
downloads is the dashboard they were looking at — plus everything needed to
defend it:

1.  **Executive Summary** — takeaways, one row per attribute, significance flags.
2.  **Segment Breakdown** — every attribute × level cell and every pairwise
    comparison, for *every* participant variable (not only the one on screen),
    with both tests' p- and q-values side by side.
3.  **Participant D-scores** — one row per recruited participant, the included
    sample first and every exclusion labelled with its reason.
4.  **Raw Trial Log** — every trial row as ingested. The audit trail.
5.  **Method & Parameters** — every setting that produced the numbers.

One engineering note. The raw log is ~175,000 rows × 23 columns for the demo
panel, and openpyxl's cell-by-cell writer takes 11–17 s for it — too slow for a
button. So the workbook is assembled with openpyxl and that one sheet's XML is
rendered column-wise with pandas string operations and streamed into the
package, which takes about a second. The header row, frozen pane, filter and
column widths still come from openpyxl; only the bulk rows are pre-rendered.
"""

from __future__ import annotations

import io
import re
import zipfile
from datetime import datetime, timezone
from xml.sax.saxutils import escape

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import get_settings
from . import charts
from .analysis import analyse_cohort, participant_table
from .ingest import TRIAL_COLUMNS
from .scoring import CohortParams

RAW_SHEET = "Raw Trial Log"
SHEETS = ["Executive Summary", "Segment Breakdown", "Participant D-scores", RAW_SHEET,
          "Method & Parameters"]

INK = "0B0B0B"
INK_2 = "52514E"
MUTED = "898781"
ACCENT = "4A3AA7"
HEAD_FILL = PatternFill("solid", fgColor="1F1D2B")
HEAD_FONT = Font(bold=True, color="FFFFFF", size=10)
BAND_FILL = PatternFill("solid", fgColor="F4F3EF")
YES_FILL = PatternFill("solid", fgColor="D5F0DD")
NO_FILL = PatternFill("solid", fgColor="F0EFEC")
WARN_FILL = PatternFill("solid", fgColor="FDF1D8")
THIN = Border(bottom=Side(style="thin", color="E1E0D9"))

FMT_D = "+0.000;-0.000;0.000"
FMT_P = '[<0.0001]"<0.0001";0.0000'
FMT_PCT = "0.0%"

METHOD_LABEL = {
    "permutation": "Permutation tests on D (distribution-free)",
    "parametric": "t-tests on log-transformed RTs",
}

PARAM_NOTES = {
    "lower_ms": "Trial-level lower trim in ms (0 = none, the Greenwald 2003 default).",
    "upper_ms": "Trial-level upper cut-off in ms (trials slower than this are removed).",
    "fast_ms": "Latency screen threshold in ms.",
    "fast_limit": "Latency screen: exclude a participant if more than this share of "
                  "combined-block trials is faster than the threshold.",
    "min_accuracy": "Accuracy screen: exclude a participant below this accuracy on the "
                    "combined blocks.",
    "min_trials": "Minimum trials per scored (role × pairing) cell.",
}


# --------------------------------------------------------------------------
# helpers
# --------------------------------------------------------------------------


def _brand(s: str | None) -> str:
    return s.title() if s and s.isupper() else (s or "")


def _yn(v) -> str:
    return "Yes" if v else "No"


def _header(ws, row: int, headers: list[str], col: int = 1) -> None:
    for j, h in enumerate(headers):
        c = ws.cell(row=row, column=col + j, value=h)
        c.fill = HEAD_FILL
        c.font = HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _write_rows(ws, start_row: int, rows: list[list], formats: dict[int, str] | None = None,
                col: int = 1) -> int:
    formats = formats or {}
    r = start_row
    for i, values in enumerate(rows):
        for j, v in enumerate(values):
            if isinstance(v, (np.floating, float)) and not np.isfinite(v):
                v = None
            c = ws.cell(row=r, column=col + j, value=v)
            if j in formats:
                c.number_format = formats[j]
            c.border = THIN
            if i % 2 == 1:
                c.fill = BAND_FILL
        r += 1
    return r


def _flag_fills(ws, ref: str) -> None:
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"Yes"'], fill=YES_FILL))
    ws.conditional_formatting.add(ref, CellIsRule(operator="equal", formula=['"No"'], fill=NO_FILL))


def _d_scale(ws, ref: str, lim: float = 0.5) -> None:
    ws.conditional_formatting.add(ref, ColorScaleRule(
        start_type="num", start_value=-lim, start_color="F6B89B",
        mid_type="num", mid_value=0, mid_color="F0EFEC",
        end_type="num", end_value=lim, end_color="9CC2EF"))


def _widths(ws, widths: dict[str, float]) -> None:
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = w


def _title(ws, text: str, sub: str | None = None) -> int:
    ws["A1"] = text
    ws["A1"].font = Font(bold=True, size=16, color=INK)
    if sub:
        ws["A2"] = sub
        ws["A2"].font = Font(size=10, color=INK_2)
    return 4


def _image(png: bytes, width_px: int) -> XLImage:
    img = XLImage(io.BytesIO(png))
    ratio = width_px / img.width
    img.width, img.height = width_px, int(img.height * ratio)
    return img


# --------------------------------------------------------------------------
# sheets
# --------------------------------------------------------------------------


def _summary(ws, batch, report: dict, method: str, generated: str) -> None:
    sample = report["sample"]
    simulated = batch.meta.get("source") == "simulated"
    row = _title(ws, "ImplicitLab — cohort topline", batch.meta.get("name", "Batch"))
    lines = [
        f"Generated {generated} · significance: {METHOD_LABEL[method]} · "
        f"Benjamini–Hochberg FDR, q < .05",
        f"Sample: {sample['included']:,} of {sample['recruited']:,} recruited participants "
        f"analysed ({sample['included_pct']:.0%}) · {sample['trial_rows']:,} trial rows ingested",
    ]
    if simulated:
        lines.append("SYNTHETIC DATA — generated for demonstration; nobody took these tests.")
    for i, text in enumerate(lines):
        c = ws.cell(row=2 + i + 1, column=1, value=text)
        c.font = Font(size=10, color=INK_2 if "SYNTHETIC" not in text else "B45309",
                      bold="SYNTHETIC" in text)
    row = 3 + len(lines) + 1

    ws.cell(row=row, column=1, value="Key takeaways").font = Font(bold=True, size=12, color=ACCENT)
    row += 1
    for t in report["takeaways"][method]:
        h = ws.cell(row=row, column=1, value=t["headline"])
        h.font = Font(bold=True, size=10, color=INK)
        h.alignment = Alignment(vertical="top", wrap_text=True)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=14)
        d = ws.cell(row=row, column=2, value=t["detail"])
        d.alignment = Alignment(vertical="top", wrap_text=True)
        d.font = Font(size=10, color=INK_2)
        ws.row_dimensions[row].height = 15 * max(2, int(len(t["detail"]) / 150) + 1)
        row += 1
    row += 1

    ws.cell(row=row, column=1, value="Key results by attribute").font = Font(bold=True, size=12, color=ACCENT)
    row += 1
    first_attr = report["attributes"][0] if report["attributes"] else {}
    headers = [
        "Attribute", "Contrast", "n", "Mean D", "95% CI low", "95% CI high",
        f"% leaning {_brand(first_attr.get('target_a'))}", "Size (IAT bands)",
        "Parametric p", "Parametric q", "Permutation p", "Permutation q",
        f"Significant (q<.05) — {method}", "Methods agree",
    ]
    _header(ws, row, headers)
    ws.freeze_panes = None
    body = []
    for a in report["attributes"]:
        o = a["overall"]
        body.append([
            a.get("pole_a") or a["key"],
            f"{_brand(a.get('target_a'))} vs {_brand(a.get('target_b'))} on "
            f"{a.get('pole_a')} vs {a.get('pole_b')}",
            o["n"], o["mean_d"], o["ci_low"], o["ci_high"], o["pct_positive"], o["magnitude"],
            o["parametric"]["p"], o["parametric"]["q"],
            o["permutation"]["p"], o["permutation"]["q"],
            _yn(o[method]["sig"]), _yn(o["agree"]),
        ])
    start = row + 1
    end = _write_rows(ws, start, body, {3: FMT_D, 4: FMT_D, 5: FMT_D, 6: FMT_PCT,
                                        8: FMT_P, 9: FMT_P, 10: FMT_P, 11: FMT_P}) - 1
    if body:
        _flag_fills(ws, f"M{start}:N{end}")
        _d_scale(ws, f"D{start}:D{end}")
    row = end + 2
    note = ws.cell(row=row, column=1, value=(
        "D > 0: faster when target A shares a key with the first attribute pole, i.e. target A "
        "is more associated with it than target B. Bands 0.15 / 0.35 / 0.65 are the Project "
        "Implicit convention, not Cohen's d. Significance flags use FDR-corrected q-values."))
    note.font = Font(italic=True, size=9, color=MUTED)
    row += 2

    ws.cell(row=row, column=1, value="Association heatmap").font = Font(bold=True, size=12, color=ACCENT)
    ws.add_image(_image(charts.heatmap(report, method), 820), f"A{row + 1}")
    row += 22
    ws.cell(row=row, column=1, value="Mean D by segment, 95% bootstrap CI").font = Font(
        bold=True, size=12, color=ACCENT)
    ws.add_image(_image(charts.forest(report, method), 820), f"A{row + 1}")

    _widths(ws, {"A": 26, "B": 38, "C": 7, "D": 10, "E": 11, "F": 11, "G": 13, "H": 14,
                 "I": 12, "J": 12, "K": 13, "L": 13, "M": 22, "N": 13})
    ws.sheet_view.showGridLines = False


def _segments(ws, batch, reports: list[dict], selected: str | None) -> None:
    row = _title(ws, "Segment breakdown",
                 "Every participant variable. q-values are corrected within each variable's "
                 "family of tests — the same families the Studio uses when that variable is "
                 "selected.")
    cell_headers = ["Attribute", "Segment", "n", "Low base (<30)", "Mean D", "SD", "Median D",
                    "95% CI low", "95% CI high", "% leaning A", "Log-RT slowdown",
                    "Parametric p", "Parametric q", "Parametric sig", "Permutation p",
                    "Permutation q", "Permutation sig", "Methods agree"]
    cmp_headers = ["Attribute", "Level A", "Level B", "n A", "n B", "Mean D (A)", "Mean D (B)",
                   "Difference", "95% CI low", "95% CI high", "Hedges' g", "Low base",
                   "Welch t", "df", "Parametric p", "Parametric q", "Parametric sig",
                   "Permutation p", "Permutation q", "Permutation sig", "Methods agree"]
    for rep in reports:
        var = rep["segment_by"]
        label = batch.meta.get("segments", {}).get(var, {}).get("label", var)
        c = ws.cell(row=row, column=1, value=f"{label}" + ("  (selected in Studio)" if var == selected else ""))
        c.font = Font(bold=True, size=13, color=ACCENT)
        row += 1
        _header(ws, row, cell_headers)
        body = []
        for a in rep["attributes"]:
            for cell in [a["overall"]] + a["levels"]:
                body.append([
                    a.get("pole_a") or a["key"], cell["level"] or "All respondents",
                    cell["n"], _yn(cell["low_base"]), cell["mean_d"], cell["sd_d"],
                    cell["median_d"], cell["ci_low"], cell["ci_high"], cell["pct_positive"],
                    cell["slowdown_pct"],
                    cell["parametric"]["p"], cell["parametric"]["q"], _yn(cell["parametric"]["sig"]),
                    cell["permutation"]["p"], cell["permutation"]["q"], _yn(cell["permutation"]["sig"]),
                    _yn(cell["agree"]),
                ])
        start = row + 1
        row = _write_rows(ws, start, body, {4: FMT_D, 5: "0.000", 6: FMT_D, 7: FMT_D, 8: FMT_D,
                                            9: FMT_PCT, 10: "+0.0%;-0.0%", 11: FMT_P, 12: FMT_P,
                                            14: FMT_P, 15: FMT_P})
        _flag_fills(ws, f"N{start}:N{row - 1}")
        _flag_fills(ws, f"Q{start}:R{row - 1}")
        _d_scale(ws, f"E{start}:E{row - 1}")
        row += 1

        if rep["comparisons"]:
            ws.cell(row=row, column=1, value="Pairwise comparisons").font = Font(bold=True, size=11, color=INK)
            row += 1
            _header(ws, row, cmp_headers)
            body = [[
                c["attribute"] if not rep["attributes"] else next(
                    (a.get("pole_a") for a in rep["attributes"] if a["key"] == c["attribute"]), c["attribute"]),
                c["level_a"], c["level_b"], c["n_a"], c["n_b"], c["mean_a"], c["mean_b"], c["diff"],
                c["ci_low"], c["ci_high"], c["hedges_g"], _yn(c["low_base"]),
                c["parametric"]["t"], c["parametric"]["df"], c["parametric"]["p"],
                c["parametric"]["q"], _yn(c["parametric"]["sig"]),
                c["permutation"]["p"], c["permutation"]["q"], _yn(c["permutation"]["sig"]),
                _yn(c["agree"]),
            ] for c in rep["comparisons"]]
            start = row + 1
            row = _write_rows(ws, start, body, {5: FMT_D, 6: FMT_D, 7: FMT_D, 8: FMT_D, 9: FMT_D,
                                                10: "0.00", 12: "0.00", 13: "0.0", 14: FMT_P,
                                                15: FMT_P, 17: FMT_P, 18: FMT_P})
            _flag_fills(ws, f"Q{start}:Q{row - 1}")
            _flag_fills(ws, f"T{start}:U{row - 1}")
        row += 2
    _widths(ws, {get_column_letter(i): 12 for i in range(1, 23)} | {"A": 16, "B": 16, "C": 16})
    ws.freeze_panes = "C4"
    ws.sheet_view.showGridLines = False


def _participants(ws, table: pd.DataFrame) -> None:
    table = table.copy()
    table["_order"] = (table["status"] != "Included").astype(int)
    table = table.sort_values(["_order", "participant_id"]).drop(columns="_order")
    cols = list(table.columns)
    ws.append(cols)
    for j in range(1, len(cols) + 1):
        c = ws.cell(row=1, column=j)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 42
    grey = Font(color=MUTED)
    for values in table.astype(object).where(table.notna(), None).itertuples(index=False, name=None):
        ws.append(list(values))
    fmts = {}
    for j, name in enumerate(cols, 1):
        if name in ("accuracy", "fast_trial_share"):
            fmts[j] = FMT_PCT
        elif name.startswith("D "):
            fmts[j] = FMT_D
        elif name.startswith("log-RT"):
            fmts[j] = "+0.0000;-0.0000"
    for j, f in fmts.items():
        for (c,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
            c.number_format = f
    excluded_from = int((table["status"] == "Included").sum()) + 2
    for r in ws.iter_rows(min_row=excluded_from, max_row=ws.max_row):
        for c in r:
            c.font = grey
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{ws.max_row}"
    _widths(ws, {get_column_letter(j): 13 for j in range(1, len(cols) + 1)}
            | {"A": 13, "B": 34})


def _raw_placeholder(ws, columns: list[str], n_rows: int) -> None:
    ws.append(columns)
    for j in range(1, len(columns) + 1):
        c = ws.cell(row=1, column=j)
        c.fill, c.font = HEAD_FILL, HEAD_FONT
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{n_rows + 1}"
    _widths(ws, {get_column_letter(j): 14 for j in range(1, len(columns) + 1)} | {"H": 22})


def _method(ws, batch, report: dict, params: CohortParams, method: str, generated: str,
            segment_by: str | None) -> None:
    row = _title(ws, "Method & parameters",
                 "Everything needed to reproduce the numbers in this workbook.")
    _header(ws, row, ["Setting", "Value", "Meaning"])
    rows = [[k, v, PARAM_NOTES.get(k, "")] for k, v in params.to_dict().items()]
    defaults = CohortParams().to_dict()
    rows += [
        ["Parameters are defaults", _yn(params.is_default()),
         "'No' means an operator changed a threshold in the Studio; defaults: "
         + ", ".join(f"{k}={v}" for k, v in defaults.items())],
        ["Significance method", METHOD_LABEL[method], "Used for the flags on the summary sheet; "
         "both methods are reported everywhere."],
        ["Segmentation on summary", segment_by or "—", ""],
    ]
    rows += [[f"Test — {k}", v, ""] for k, v in (report.get("method_note") or {}).items()]
    rows += [["Interpretation bands", report.get("interpretation_note", ""), ""]]
    ing = batch.report or {}
    rows += [
        ["Batch ID", batch.id, ""],
        ["Batch name", batch.meta.get("name"), ""],
        ["Source", batch.meta.get("source"), batch.meta.get("note", "")],
        ["Seed", batch.meta.get("seed", "—"), "Simulated batches regenerate exactly from this seed."],
        ["File", batch.meta.get("filename", "—"), ""],
        ["Rows read / accepted / rejected",
         f"{ing.get('rows_read', '—')} / {ing.get('rows_accepted', '—')} / {ing.get('rows_rejected', '—')}",
         "; ".join(f"{k}: {v}" for k, v in (ing.get("rejection_reasons") or {}).items())],
        ["Duplicates removed", ing.get("duplicates_removed", 0), ""],
        ["App version", get_settings().version, "ImplicitLab"],
        ["Generated", generated, ""],
        ["Reproduce", f"/api/cohort/batches/{batch.id}/export.xlsx?"
         + "&".join(f"{k}={v}" for k, v in params.to_dict().items())
         + f"&method={method}" + (f"&segment_by={segment_by}" if segment_by else ""),
         "Every analysis is a pure function of (batch, parameters); resampling is seeded."],
    ]
    end = _write_rows(ws, row + 1, rows)
    for r in ws.iter_rows(min_row=row + 1, max_row=end, max_col=3):
        for c in r:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    _widths(ws, {"A": 30, "B": 48, "C": 80})
    ws.sheet_view.showGridLines = False


# --------------------------------------------------------------------------
# raw sheet streaming
# --------------------------------------------------------------------------

_ILLEGAL = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _col_xml(s: pd.Series) -> pd.Series:
    """Render one column as <c> elements, vectorised."""
    if pd.api.types.is_bool_dtype(s):
        return np.where(s.to_numpy(dtype=bool), '<c t="b"><v>1</v></c>', '<c t="b"><v>0</v></c>')
    if pd.api.types.is_numeric_dtype(s) and not isinstance(s.dtype, pd.CategoricalDtype):
        v = pd.to_numeric(s, errors="coerce")
        txt = v.map(lambda x: format(x, ".10g"))
        return np.where(v.notna(), "<c><v>" + txt + "</v></c>", "<c/>")
    txt = s.astype(object).where(s.notna(), "").astype(str)
    uniq = pd.unique(txt)
    rendered = {u: (f'<c t="inlineStr"><is><t>{escape(_ILLEGAL.sub("", u))}</t></is></c>'
                    if u else "<c/>") for u in uniq}
    return txt.map(rendered).to_numpy()


RAW_CHUNK = 5000


def _raw_rows_xml(df: pd.DataFrame, start_row: int = 2) -> str:
    parts = [_col_xml(df[c]) for c in df.columns]
    body = parts[0].astype(object)
    for p in parts[1:]:
        body = body + p
    idx = np.arange(start_row, start_row + len(df)).astype(str)
    rows = '<row r="' + idx.astype(object) + '">' + body + "</row>"
    return "".join(rows.tolist())


def _inject_raw(xlsx: bytes, sheet_index: int, raw: pd.DataFrame, ref: str) -> bytes:
    """Stream the raw rows into the placeholder sheet, one chunk at a time.

    Rendering all 175k rows as one string holds several million Python string
    objects at once — over a gigabyte, which is more than a free-tier instance
    has. Written straight into the zip entry in 5,000-row chunks, peak memory
    is bounded by the chunk, not the batch.
    """
    src = zipfile.ZipFile(io.BytesIO(xlsx))
    target = f"xl/worksheets/sheet{sheet_index}.xml"
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED, compresslevel=5) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename != target:
                dst.writestr(item, data)
                continue
            xml = data.decode("utf-8")
            xml = re.sub(r'<dimension ref="[^"]*"', f'<dimension ref="{ref}"', xml, count=1)
            end = xml.index("</sheetData>")
            info = zipfile.ZipInfo(item.filename, date_time=item.date_time)
            info.compress_type = zipfile.ZIP_DEFLATED
            with dst.open(info, "w", force_zip64=True) as fh:
                fh.write(xml[:end].encode("utf-8"))
                for start in range(0, len(raw), RAW_CHUNK):
                    chunk = raw.iloc[start:start + RAW_CHUNK]
                    fh.write(_raw_rows_xml(chunk, start_row=start + 2).encode("utf-8"))
                fh.write(xml[end:].encode("utf-8"))
    return out.getvalue()


# --------------------------------------------------------------------------
# entry point
# --------------------------------------------------------------------------


def build_workbook(batch, params: CohortParams, *, segment_by: str | None = None,
                   method: str = "permutation") -> bytes:
    generated = datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")
    report = analyse_cohort(batch.prep, batch.meta, params, segment_by=segment_by)
    selected = report["segment_by"]
    others = [
        analyse_cohort(batch.prep, batch.meta, params, segment_by=v, include_distributions=False)
        for v in batch.meta.get("segments", {}) if v != selected
    ]
    table, _ = participant_table(batch.prep, batch.meta, params)

    segs = list(batch.meta.get("segments", {}))
    raw = batch.trials
    raw_cols = ["participant_id", *segs] + [c for c in TRIAL_COLUMNS
                                            if c in raw.columns and c not in ("participant_id", *segs)]
    raw = raw[raw_cols]

    wb = Workbook()
    ws = wb.active
    ws.title = SHEETS[0]
    _summary(ws, batch, report, method, generated)
    _segments(wb.create_sheet(SHEETS[1]), batch, [report] + others, selected)
    _participants(wb.create_sheet(SHEETS[2]), table)
    _raw_placeholder(wb.create_sheet(RAW_SHEET), raw_cols, len(raw))
    _method(wb.create_sheet(SHEETS[4]), batch, report, params, method, generated, selected)

    buf = io.BytesIO()
    wb.save(buf)
    ref = f"A1:{get_column_letter(len(raw_cols))}{len(raw) + 1}"
    return _inject_raw(buf.getvalue(), SHEETS.index(RAW_SHEET) + 1, raw, ref)
